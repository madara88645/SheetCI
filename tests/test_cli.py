import json
from pathlib import Path
from typer.testing import CliRunner
from sheetci.cli import app

runner = CliRunner()

def test_scan_broken_model():
    # Scanning broken model should return exit code 1
    result = runner.invoke(app, ["scan", "examples/broken-commission-model.xlsx"])
    assert result.exit_code == 1
    assert "SheetCI Scan Results" in result.stdout
    assert "Risk Score:          100/100" in result.stdout
    assert "Status:              FAIL" in result.stdout

def test_scan_clean_model():
    # Scanning clean model should return exit code 0
    result = runner.invoke(app, ["scan", "examples/clean-model.xlsx"])
    assert result.exit_code == 0
    assert "SheetCI Scan Results" in result.stdout
    assert "Risk Score:          0/100" in result.stdout
    assert "Status:              PASS" in result.stdout

def test_scan_json_output():
    result = runner.invoke(app, ["scan", "examples/broken-commission-model.xlsx", "--json"])
    assert result.exit_code == 1
    # Verify valid JSON
    data = json.loads(result.stdout)
    assert data["metadata"]["workbook_name"] == "broken-commission-model.xlsx"
    assert data["metadata"]["status"] == "FAIL"
    assert len(data["findings"]) == 8

def test_reports_generation(tmp_path):
    md_report = tmp_path / "report.md"
    html_report = tmp_path / "report.html"
    
    result = runner.invoke(app, [
        "scan", 
        "examples/broken-commission-model.xlsx", 
        "--out", str(md_report),
        "--html", str(html_report)
    ])
    
    assert result.exit_code == 1
    assert md_report.exists()
    assert html_report.exists()
    
    md_content = md_report.read_text(encoding="utf-8")
    assert "# SheetCI Audit Report" in md_content
    assert "**Status**: **FAIL**" in md_content
    
    html_content = html_report.read_text(encoding="utf-8")
    assert "<title>SheetCI Audit Report" in html_content
    assert 'id="sheetci-report"' in html_content

def test_console_output_shows_occurrence_counts(tmp_path):
    import openpyxl

    wb_path = tmp_path / "repeated.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    for r in range(2, 12):
        ws[f"A{r}"] = r
        ws[f"B{r}"] = f"=A{r}*7.5"
    wb.save(wb_path)

    result = runner.invoke(app, ["scan", str(wb_path)])

    assert result.exit_code == 0
    assert "B2:B11" in result.stdout
    assert "x10" in result.stdout


def test_markdown_report_shows_occurrence_counts(tmp_path):
    md_report = tmp_path / "report.md"
    result = runner.invoke(app, [
        "scan",
        "examples/broken-commission-model.xlsx",
        "--out", str(md_report),
    ])
    assert result.exit_code == 1
    content = md_report.read_text(encoding="utf-8")
    assert "**Location**:" in content
    assert "Occurrences" in content


def test_realistic_model_passes():
    """A healthy workbook must not fail. This is the regression that started this work."""
    result = runner.invoke(app, ["scan", "examples/realistic-model.xlsx", "--json"])
    assert result.exit_code == 0
    data = json.loads(result.stdout)
    assert data["metadata"]["status"] == "PASS"
    assert data["metadata"]["counts"]["critical"] == 0
    assert data["metadata"]["counts"]["warning"] == 0
    assert len(data["findings"]) <= 5


def test_broken_model_still_detects_criticals():
    """Noise reduction must not cost detection power."""
    result = runner.invoke(app, ["scan", "examples/broken-commission-model.xlsx", "--json"])
    assert result.exit_code == 1
    data = json.loads(result.stdout)
    assert data["metadata"]["status"] == "FAIL"
    rules = {f["rule_id"] for f in data["findings"]}
    assert "BROKEN_REF" in rules
    assert "SELF_REFERENCE" in rules
    assert "CACHED_ERROR" in rules
